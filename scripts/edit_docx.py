from docx import Document 
from pathlib import Path
doc = Document("HILLARY_SCHEDULE APRIL.docx")
aliases_names = {
    "Iannaccio": "Pascoli", 
    "Crescenza": "Pignatelli", 
    "Carmentano": "Gallo",
    "Bettoni": "Ferrari",
    "Arca": "Greco",
    "Crippa": "Conti",
    "Tasso": "De Luca",
    "Cirocco": "Gallo",
    "Finardi": "Ricci",
    "Betta": "Rossi",
    "Facheris": "Bianchi",
    "Pandini": "Galli",
    "Accarino": "Esposito",
    "Merletti": "Neri",
    "Cirillo": "Verdi",
    "SIA" : "TEC",
    "ITE": "ETI",
    "AFM": "MAF",
     "LSU": "LSS",
     "LES": "LSE"
}
script_dir = Path(__file__).parent
try: 
    for file_path in script_dir.glob("*.docx"): 
        doc = Document(file_path)
        for table in doc.tables: 
            for row in table.rows:
                for cell in row.cells:
                    for key, value in aliases_names.items():
                        if key in cell.text:
                            cell.text = cell.text.replace(key, value)
        doc.save(file_path)
except Exception as e:
    print(f"An error occurred: {e}")


